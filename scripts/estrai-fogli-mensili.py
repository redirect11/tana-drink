# -*- coding: utf-8 -*-
# I FOGLI MENSILI DI ANALISI DATI.xlsx IN UN FILE JSON (REQ-CASSA-011).
#
# Ogni mese ha un foglio suo («GEN 25», «GIU26», …), diviso in settimane
# come nel Bilancio dell'app: blocchi di sette giorni dal primo del mese
# (I SETT, II SETT, III SETT, IV SETT) e i giorni che restano (EXTRA). Per
# ogni giorno tre righe — MINIMO, REALE, RAPP. — e sopra, a volte, una nota
# (KARAOKE, GIGLI, CENE…). In fondo al mese STIPENDI, SPESE e INCASSI, scritti
# negativi i primi due.
#
# Lo legge Python perché fra le dipendenze del progetto non c'è un lettore di
# xlsx; a scrivere nel database è scripts/importa-fogli-mensili.mjs.
#
#   python scripts/estrai-fogli-mensili.py "ANALISI DATI.xlsx" backup/fogli-mensili.json

import io
import json
import re
import sys

import openpyxl

MESI = {'GEN': 1, 'FEB': 2, 'MAR': 3, 'APR': 4, 'MAG': 5, 'GIU': 6,
        'LUG': 7, 'AGO': 8, 'SET': 9, 'OTT': 10, 'NOV': 11, 'DIC': 12}

# I tre blocchi di settimane: riga della nota, della data, del minimo e del
# reale; e le due metà di colonne (C..I e L..R).
BLOCCHI = [(3, 5, 6, 7), (11, 13, 14, 15), (19, 21, 22, 23)]
COLONNE = [list('CDEFGHI'), list('LMNOPQR')]


def numero(v):
    return float(v) if isinstance(v, (int, float)) else None


def mese_del_foglio(nome):
    m = re.match(r'^\s*([A-Z]{3})\s*(\d{2})\s*$', nome)
    if not m or m.group(1) not in MESI:
        return None
    return '20%s-%02d' % (m.group(2), MESI[m.group(1)])


def leggi(percorso):
    wb = openpyxl.load_workbook(percorso, data_only=True)
    mesi = []
    for ws in wb.worksheets:
        mese = mese_del_foglio(ws.title)
        if not mese:
            continue
        giorni = {}
        for (r_nota, r_data, r_min, r_reale) in BLOCCHI:
            for metà in COLONNE:
                for col in metà:
                    g = numero(ws['%s%d' % (col, r_data)].value)
                    if g is None or not (1 <= g <= 31):
                        continue
                    nota = ws['%s%d' % (col, r_nota)].value
                    giorni['%s-%02d' % (mese, int(g))] = {
                        'minimo': numero(ws['%s%d' % (col, r_min)].value),
                        # Il REALE e gli INCASSI del foglio l'app non li
                        # mostra (l'incassato sono i conti): restano per
                        # confrontarli, perché non sempre tornano.
                        'reale': numero(ws['%s%d' % (col, r_reale)].value),
                        'nota': str(nota).strip() if nota not in (None, '') else None,
                    }
        # In fondo al mese: le etichette in Q, i numeri in S.
        fondo = {}
        for r in range(18, 30):
            etichetta = ws['Q%d' % r].value
            if isinstance(etichetta, str):
                fondo[etichetta.strip().upper()] = numero(ws['S%d' % r].value)
        mesi.append({
            'mese': mese,
            'foglio': ws.title,
            'giorni': giorni,
            # Nel foglio stipendi e spese sono negativi: qui sono quanto è uscito.
            'stipendi': abs(fondo['STIPENDI']) if fondo.get('STIPENDI') is not None else None,
            'spese': abs(fondo['SPESE']) if fondo.get('SPESE') is not None else None,
            'incassi': fondo.get('INCASSI'),
        })
    return sorted(mesi, key=lambda m: m['mese'])


if __name__ == '__main__':
    sorgente, uscita = sys.argv[1], sys.argv[2]
    mesi = leggi(sorgente)
    with io.open(uscita, 'w', encoding='utf-8') as f:
        f.write(json.dumps(mesi, ensure_ascii=False, indent=1))
    for m in mesi:
        print('%s  %-7s %2d giorni  stipendi %-8s spese %-8s incassi %s' % (
            m['mese'], m['foglio'], len(m['giorni']), m['stipendi'], m['spese'], m['incassi']))
    print('scritto %s: %d mesi' % (uscita, len(mesi)))
